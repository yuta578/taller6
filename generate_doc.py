import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def generate_report():
    doc = docx.Document()

    # Configuración de márgenes
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Estilo base
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    def add_header(title, subtitle):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title)
        run.font.size = Pt(22)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run2 = p2.add_run(subtitle)
        run2.font.size = Pt(12)
        run2.font.italic = True
        run2.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
        doc.add_paragraph()

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

    def add_callout(text, title=''):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        shading = parse_xml(r'<w:shd {} w:fill="F3F4F6"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shading)
        
        borders = parse_xml(r'''
            <w:tcBorders {} >
                <w:top w:val="none"/>
                <w:left w:val="single" w:sz="24" w:space="0" w:color="3B82F6"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        '''.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(borders)
        
        cp = cell.paragraphs[0]
        cp.paragraph_format.space_before = Pt(6)
        cp.paragraph_format.space_after = Pt(6)
        if title:
            r_t = cp.add_run(title + '\n')
            r_t.bold = True
            r_t.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        r_body = cp.add_run(text)
        r_body.font.size = Pt(10.5)
        doc.add_paragraph()

    # Encabezado
    add_header(
        'TALLER 6: CEREMONIAS ÁGILES EN CONTINUOUS DEPLOYMENT',
        'Resolución Integral del Caso: “La funcionalidad de Pagos Express ya está en main, pero nadie se atreve a activarla”\nEntorno Operativo: Render + GitHub Actions + ConfigCat + Scrum & TBD Playbook v3.0'
    )

    # 0. Apertura
    add_h1('0. Apertura y Conexión con el Nuevo Flujo')
    add_callout(
        '“Ya tenemos el pipeline completo: integramos a main, se despliega solo a Render y protegemos funcionalidades con ConfigCat.\n'
        '¿Cómo cambian el Sprint Review y la Retrospective cuando desplegamos varias veces por semana?”\n\n'
        '1. El Sprint Review pasa de ser una aprobación tardía de fin de sprint a una demostración de valor real en producción con Feature Toggles y telemetría de negocio.\n'
        '2. La Sprint Retrospective pasa de ser solo una conversación subjetiva a una auditoría objetiva de la salud del pipeline (CI/CD), la disciplina TBD y el flujo de valor (métricas DORA).',
        'Pregunta Clave del Taller'
    )

    add_h2('¿Qué parte del Review o de la Retro actual se siente más ‘antigua’ o poco útil?')
    doc.add_paragraph('• En el Review clásico: Se siente obsoleto montar una demo en localhost o diapositivas para mostrar código que ya lleva 9 días integrado en main y desplegado en Render. Mostrar "qué construimos" en vez de "qué valor recibieron los usuarios" no aporta al negocio.')
    doc.add_paragraph('• En la Retro clásica: Hablar de estimaciones abstractas ignorando que el CI estuvo roto 4 días por un test flaky, que hubo una caída de 11 minutos por un health check mal configurado, o que hay 3 feature toggles con más de 45 días acumulando deuda técnica.')

    # 1. Sprint Review
    add_h1('1. Adaptación del Sprint Review al Continuous Deployment')
    add_h2('1.1 Principios del Nuevo Review: Clásico vs CD + TBD')

    table1 = doc.add_table(rows=5, cols=3)
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ['Dimensión', 'Sprint Review Clásico', 'Sprint Review en CD + TBD (Playbook 3.0)']
    for i, h in enumerate(headers):
        cell = table1.cell(0, i)
        cell.paragraphs[0].add_run(h).bold = True
        shading = parse_xml(r'<w:shd {} w:fill="1E3A8A"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shading)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    data1 = [
        ('Ambiente de Demo', 'Staging o Localhost ("en mi máquina funciona").', 'Producción real en Render con datos vivos.'),
        ('Naturaleza', 'Todo el incremento junto ("Todo o nada").', 'Atómico y desacoplado mediante Feature Toggles.'),
        ('Feedback', 'Tardío, acumulado cada 14 días.', 'Continuo con métricas de cohortes en tiempo real.'),
        ('Enfoque', 'Demostrar qué código se escribió.', 'Demostrar qué valor recibieron los usuarios.')
    ]
    for row_idx, row_data in enumerate(data1, start=1):
        for col_idx, text in enumerate(row_data):
            cell = table1.cell(row_idx, col_idx)
            cell.paragraphs[0].add_run(text)
            if row_idx % 2 == 1:
                shading = parse_xml(r'<w:shd {} w:fill="F9FAFB"/>'.format(nsdecls('w')))
                cell._tc.get_or_add_tcPr().append(shading)

    doc.add_paragraph()

    add_h2('1.2 Simulación de Sprint Review: Caso Pagos Express')
    add_callout(
        '1. ¿Qué problema de usuario resuelve?\n'
        '   El checkout tradicional de 4 pasos tiene un 68% de abandono de carrito al exigir teclear 16 dígitos de tarjeta. Pagos Express permite pagar en 1 clic con tarjeta tokenizada guardada.\n\n'
        '2. ¿Está activo para todos o detrás de toggle?\n'
        '   Está integrado en main y desplegado en Render hace 9 días. Está protegido 100% por el toggle pagos-express-v1 en ConfigCat (inicialmente en OFF).\n\n'
        '3. Demo en vivo en Render:\n'
        '   • Con Toggle OFF: Se muestra el formulario tradicional con alta fricción y latencia de 420 ms.\n'
        '   • Con Toggle ON: La interfaz se transforma mostrando la tarjeta Visa 4242 tokenizada y el botón de 1 clic, procesando la compra en 85 ms.\n\n'
        '4. Métrica o evidencia de valor:\n'
        '   Latencia reducida a 85 ms, 0% de errores en pasarela y proyección de +93% de incremento en conversión.\n\n'
        '5. Pregunta al Stakeholder:\n'
        '   “¿Esto cumple con la expectativa para habilitar el Canary del 20% este viernes?” (Aprobado).',
        'Estructura de la Demostración (4 minutos)'
    )

    # 2. Retrospective
    add_h1('2. Adaptación de la Sprint Retrospective')
    add_h2('2.1 Diagnóstico de los 7 Hechos Problemáticos del Caso')
    doc.add_paragraph('1. Presión del PO por activar el 20% el viernes: Resuelto mediante Progressive Delivery con Canary controlado y observabilidad activa.')
    doc.add_paragraph('2. Último deploy falló health check 11 minutos: Resuelto implementando endpoint /api/health dedicado en Render.')
    doc.add_paragraph('3. 3 Toggles llevan > 45 días activos: Resuelto formalizando TTL máximo de 30 días y generando el PR de refactorización #142.')
    doc.add_paragraph('4. Solo 4 integraciones a main en el sprint: Resuelto restaurando la disciplina de Trunk-Based Development con ramas < 24h.')
    doc.add_paragraph('5. Rama viva 4 días por test flaky: Resuelto erradicando pruebas inestables y blindando CI con 24 tests deterministas.')
    doc.add_paragraph('6. Ceguera de métricas de pago: Resuelto actualizando el DoD v3.0 para exigir telemetría de errores y latencia antes del release.')
    doc.add_paragraph('7. Stakeholder anunció que ya está listo: Resuelto educando en la distinción Deploy != Release.')

    add_h2('2.2 Matriz de Retrospective y 3 Acciones Concretas')
    table2 = doc.add_table(rows=2, cols=2)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    t2_data = [
        ('🟢 LO QUE FUNCIONÓ BIEN\n• Código sliceado e integrado con toggle en main.\n• ConfigCat protegió la producción durante 9 días.\n• Pipeline de Render desplegando sin intervención manual.',
         '🟡 LO QUE GENERÓ FRICCIÓN\n• Marketing lanzó campaña antes de validación.\n• Miedo a encender el toggle por falta de métricas.\n• 3 toggles viejos ensuciando el código.'),
        ('🔴 LO QUE ROMPIÓ EL FLUJO O GENERÓ MIEDO\n• Test flaky tuvo CI rojo 4 días (rama viva 4 días).\n• Caída de 11 minutos por health check fallido.\n• Miedo del equipo por falta de telemetría de pagos.',
         '🚀 ACCIONES DE MEJORA (MÁXIMO 3)\n1. Erradicar test flaky y fijar /api/health en Render (Carlos | 48h).\n2. Exigir SLIs de pago en el DoD v3.0 (Sofía | Viernes).\n3. Ejecutar PR #142 de limpieza de toggles > 30 días (Laura + Carlos | Sprint+1).')
    ]
    for r_idx, (c1, c2) in enumerate(t2_data):
        cell1 = table2.cell(r_idx, 0)
        cell2 = table2.cell(r_idx, 1)
        cell1.paragraphs[0].add_run(c1)
        cell2.paragraphs[0].add_run(c2)
        shading1 = parse_xml(r'<w:shd {} w:fill="F3F4F6"/>'.format(nsdecls('w')))
        shading2 = parse_xml(r'<w:shd {} w:fill="F9FAFB"/>'.format(nsdecls('w')))
        cell1._tc.get_or_add_tcPr().append(shading1)
        cell2._tc.get_or_add_tcPr().append(shading2)

    doc.add_paragraph()

    # 3. Liberación Segura
    add_h1('3. Liberación Segura de Valor (Progressive Delivery)')
    add_callout(
        '• FASE 0: Dark Deploy (0% usuarios) — Código en producción en OFF.\n'
        '• FASE 1: Dogfooding Interno — Solo cuentas @nexusdash.dev para pruebas de humo.\n'
        '• FASE 2: Canary 20% (Viernes) — Hashing por userId, monitoreando tasa de error < 0.5%.\n'
        '• FASE 3: 100% General Availability (Lunes) — Liberación completa si las métricas se mantienen estables.\n\n'
        '🚨 PROTOCOLO KILL SWITCH:\n'
        'Si la tasa de error supera el 1.0% o la latencia excede 500 ms, se apaga el toggle a 0% en segundos desde ConfigCat sin requerir nuevo deploy ni reinicio de servicios.',
        'Estrategia de Liberación y Criterios Go / No-Go'
    )

    # 4. Playbook y Cierre
    add_h1('4. Actualización del Playbook y Checklist de Fase 1')
    doc.add_paragraph('• Scrum + TBD Playbook v3.0 formalizado con las nuevas ceremonias adaptadas.')
    doc.add_paragraph('• Definition of Done (DoD) v3.0 integrado con criterios técnicos, despliegue, toggles y observabilidad.')
    doc.add_paragraph('• Checklist de Fase 1 cerrado al 100% con éxito.')
    add_callout(
        '“A partir de hoy, nunca más lanzaremos una campaña de marketing asumiendo que deploy = listo; coordinaremos los porcentajes de rollout en ConfigCat observando datos reales y no toleraremos pruebas inestables en el CI.”',
        'Ronda Final del Taller'
    )

    # Guardar documento
    doc.save(r'c:\Users\ANGEL\Desktop\taller6\Taller 6.docx')
    doc.save(r'c:\Users\ANGEL\Desktop\Taller 6.docx')
    print('Taller 6.docx creado con éxito.')

if __name__ == '__main__':
    generate_report()
